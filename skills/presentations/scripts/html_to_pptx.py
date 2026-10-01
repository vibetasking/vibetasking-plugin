import asyncio
import importlib.util
import os
import sys
import tempfile
from pathlib import Path

from playwright.async_api import async_playwright
from pptx import Presentation
from pptx.util import Inches

# Reuse the documents skill's shared Chromium launcher (bundled Chromium, then
# system Chrome fallback) so a missing bundled browser doesn't break export.
_PDF_LIB_PATH = Path(__file__).resolve().parent.parent.parent / "documents" / "scripts" / "_pdf_lib.py"
_spec = importlib.util.spec_from_file_location("_pdf_lib_shared", _PDF_LIB_PATH)
_pdf_lib = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(_pdf_lib)
launch_chromium = _pdf_lib.launch_chromium

# A4 landscape in inches
_SLIDE_WIDTH = Inches(11.6929)
_SLIDE_HEIGHT = Inches(8.2677)

# Viewport at 96 DPI for A4 landscape (297mm x 210mm)
_VIEWPORT_WIDTH = 1122
_VIEWPORT_HEIGHT = 793

# 2x for retina-quality screenshots
_DEVICE_SCALE_FACTOR = 2


async def html_slides_to_pptx(html_path: str, output_path: str) -> str:
    """Screenshot each .page div via Playwright and assemble the images into a PPTX."""

    html_path = os.path.abspath(html_path)
    output_path = os.path.abspath(output_path)

    prs = Presentation()
    prs.slide_width = _SLIDE_WIDTH
    prs.slide_height = _SLIDE_HEIGHT
    blank_layout = prs.slide_layouts[6]  # Blank slide

    tmp_dir = tempfile.mkdtemp(prefix="html_to_pptx_")

    try:
        async with async_playwright() as p:
            browser = await launch_chromium(p, headless=True)
            context = await browser.new_context(
                viewport={"width": _VIEWPORT_WIDTH, "height": _VIEWPORT_HEIGHT},
                device_scale_factor=_DEVICE_SCALE_FACTOR,
            )
            page = await context.new_page()

            await page.goto(
                f"file://{html_path}",
                wait_until="networkidle",
                timeout=60000,
            )

            slides = await page.query_selector_all(".page")
            if not slides:
                raise ValueError("No elements with class 'page' found in the HTML.")

            for i, element in enumerate(slides):
                png_path = os.path.join(tmp_dir, f"slide_{i:03d}.png")
                await element.screenshot(path=png_path, type="png", animations="disabled")

                slide = prs.slides.add_slide(blank_layout)
                slide.shapes.add_picture(
                    png_path,
                    left=0,
                    top=0,
                    width=prs.slide_width,
                    height=prs.slide_height,
                )

            await browser.close()

        prs.save(output_path)
        return output_path

    finally:
        for f in Path(tmp_dir).glob("*.png"):
            f.unlink(missing_ok=True)
        Path(tmp_dir).rmdir()


def main():
    if len(sys.argv) != 3:
        print(f"Usage: python {sys.argv[0]} <input.html> <output.pptx>")
        sys.exit(1)

    html_path, output_path = sys.argv[1], sys.argv[2]

    if not os.path.isfile(html_path):
        print(f"Error: {html_path} not found")
        sys.exit(1)

    asyncio.run(html_slides_to_pptx(html_path, output_path))
    print(f"Created {output_path}")


if __name__ == "__main__":
    main()
